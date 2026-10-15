from pathlib import Path
import subprocess
from pypdf import PdfReader, PdfWriter
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'output/docx'
TMP=ROOT/'tmp/mobilenet'
OUT.mkdir(parents=True,exist_ok=True)
TMP.mkdir(parents=True,exist_ok=True)
page=PdfReader(r'C:\Users\bash\Desktop\1704.04861v1.pdf').pages[3]
page.cropbox.lower_left=(73,602)
page.cropbox.upper_right=(263,720)
w=PdfWriter(); w.add_page(page)
with open(TMP/'source_figure.pdf','wb') as f: w.write(f)
subprocess.run(['pdftoppm','-cropbox','-r','220','-png','-singlefile',str(TMP/'source_figure.pdf'),str(TMP/'source_figure')],check=True)
d=Document()
s=d.sections[0]
s.page_width=Inches(8.5); s.page_height=Inches(11)
s.top_margin=s.bottom_margin=Inches(.8)
s.left_margin=s.right_margin=Inches(.85)
for name in ['Normal','Title','Heading 1','Heading 2','Caption']:
    st=d.styles[name]; st.font.name='Times New Roman'; st.font.color.rgb=RGBColor(0,0,0)
    st.font.size=Pt(11.5); st.font.bold=False
    st.paragraph_format.space_after=Pt(7)
    st.paragraph_format.line_spacing=1.12
d.styles['Title'].font.size=Pt(16)
d.styles['Heading 1'].font.size=Pt(13)
d.styles['Heading 2'].font.underline=True
for name in ['Heading 1','Heading 2']:
    d.styles[name].paragraph_format.keep_with_next=True
    d.styles[name].paragraph_format.space_before=Pt(9)
d.styles['Normal'].paragraph_format.widow_control=True
d.styles['Caption'].font.size=Pt(10)
d.styles['Caption'].font.italic=True

def p(t):
    x=d.add_paragraph(t); x.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    return x
def h(t): d.add_paragraph(t,'Heading 2')
def new(): d.add_page_break()
def eq(tokens):
    para=d.add_paragraph(); para.alignment=WD_ALIGN_PARAGRAPH.CENTER
    m=OxmlElement('m:oMath')
    for text,sub,sup in tokens:
        base=OxmlElement('m:r'); t=OxmlElement('m:t'); t.text=text; base.append(t)
        if sub:
            outer=OxmlElement('m:sSub'); e=OxmlElement('m:e'); e.append(base); outer.append(e)
            su=OxmlElement('m:sub'); r=OxmlElement('m:r'); t=OxmlElement('m:t'); t.text=sub; r.append(t); su.append(r); outer.append(su); base=outer
        if sup:
            outer=OxmlElement('m:sSup'); e=OxmlElement('m:e'); e.append(base); outer.append(e)
            su=OxmlElement('m:sup'); r=OxmlElement('m:r'); t=OxmlElement('m:t'); t.text=sup; r.append(t); su.append(r); outer.append(su); base=outer
        m.append(base)
    para._p.append(m)

d.add_paragraph('Architecture 4 MobileNetV1','Title')
h('Introduction')
p('MobileNet is a lightweight convolutional neural network introduced by Howard et al. for mobile and embedded devices [4]. These devices have limited memory and processing power, so the model needs to be accurate while keeping computation low. MobileNet does this by splitting a standard convolution into two simpler operations. For our facial recognition system, this is useful because we need to consider how quickly a face can be processed along with how accurately it can be recognised.')
h('Depthwise separable convolution')
p('In a standard convolution, filtering the image features and combining information from different channels happen in the same operation. MobileNet splits this into two steps. A depthwise convolution applies one filter to each input channel separately. A pointwise convolution then uses 1 × 1 filters to combine these outputs and produce the required number of output channels [4].')
p('For example, consider a 224 × 224 × 3 input and 64 output channels, using a 3 × 3 kernel with stride 1 and padding that preserves the image size. A standard convolution uses 64 filters of size 3 × 3 × 3 and produces a 224 × 224 × 64 output. The depthwise step instead uses three 3 × 3 filters and produces a 224 × 224 × 3 output. The pointwise step then uses 64 filters of size 1 × 1 × 3 to produce the same output dimensions. This illustrates the operation; the actual first layer of MobileNet uses a standard convolution with stride 2 [4].')
x=d.add_paragraph(); x.alignment=WD_ALIGN_PARAGRAPH.CENTER
x.add_run().add_picture(str(TMP/'source_figure.png'),width=Inches(3.25))
x=d.add_paragraph('Figure 1. Standard convolution and depthwise separable convolution with batch normalization and ReLU. Reproduced from Figure 3 of Howard et al. [4].','Caption'); x.alignment=WD_ALIGN_PARAGRAPH.CENTER

new(); h('Computational cost and parameters')
p('In the equations below, D with subscript K is the kernel width and height, and D with subscript F is the width and height of the feature map. M and N are the numbers of input and output channels. The equations assume stride 1 and equal input and output spatial dimensions, as in the original paper [4]. Computation is measured in multiply-adds using the paper’s convention.')
eq([('C','standard',None),(' = ',None,None),('D','K','2'),('MN',None,None),('D','F','2')])
eq([('C','separable',None),(' = ',None,None),('D','K','2'),('M',None,None),('D','F','2'),(' + MN',None,None),('D','F','2')])
p('The first term in the second equation is the depthwise cost, and the second is the pointwise cost. Dividing the separable cost by the standard cost gives the ratio below. For 3 × 3 kernels, the paper reports roughly eight to nine times less computation [4]. The actual saving depends on the number of output channels, and the speed on a device also depends on how the operations are implemented.')
eq([('C','separable',None),(' / ',None,None),('C','standard',None),(' = 1/N + 1/',None,None),('D','K','2')])
p('The weight counts follow the same pattern, with the spatial feature-map factor removed. Standard convolution needs one full filter for each output channel. Depthwise separable convolution needs one spatial filter per input channel, followed by the pointwise weights. In the example above, this reduces the count from 1,728 weights to 219, excluding biases and batch-normalization parameters. Computation falls from 86,704,128 to 10,988,544 multiply-adds, or about 7.9 times less. These values are calculated from the paper’s equations.')
p('For the complete baseline ImageNet model, Howard et al. report 4.2 million parameters and 569 million multiply-adds, with 70.6% classification accuracy. Their comparable model using full convolutions requires 29.3 million parameters and 4,866 million multiply-adds and reaches 71.7%. The reduction in cost therefore comes with a 1.1 percentage-point decrease in ImageNet accuracy [4, Table 4]. These are ImageNet results, so they should not be presented as face-verification results.')
h('Use of GEMM')
p('GEMM means general matrix multiplication. MobileNet uses this operation to carry out its pointwise convolutions efficiently. A standard convolution often needs the input data to be rearranged in memory through a step called im2col before GEMM can be applied. A 1 × 1 pointwise convolution can use GEMM directly, which avoids this extra step [4].')
p('This matters because pointwise convolutions account for 94.86% of MobileNet’s multiply-adds and 74.59% of its parameters in the original analysis. Most of the remaining parameters are in the fully connected layer [4, Table 2]. The architecture therefore places most of its work in an operation that can use highly optimized matrix-multiplication routines.')

new(); h('Architecture flow')
p('The baseline MobileNet takes a 224 × 224 RGB image. Its first layer is a full 3 × 3 convolution, followed by blocks containing a 3 × 3 depthwise convolution and a 1 × 1 pointwise convolution. The number of channels increases through 32, 64, 128, 256, 512 and 1,024 as the spatial dimensions are reduced. Five repeated blocks operate at the 14 × 14 × 512 stage. The final feature map is 7 × 7 × 1,024 [4, Table 1].')
p('Batch normalization and ReLU follow the convolutional layers, including both parts of each depthwise separable block. ReLU introduces non-linearity, allowing the network to learn more than a sequence of linear transformations. Average pooling then reduces the final feature map to a 1 × 1 × 1,024 representation. A fully connected layer produces 1,000 class scores for ImageNet, followed by softmax. The final fully connected layer has no ReLU. Counting depthwise and pointwise convolutions separately, the network has 28 layers [4].')
h('Width multiplier')
p('The width multiplier, α, makes the network thinner by reducing the number of channels. For a layer with M input channels and N output channels, the reduced layer uses αM and αN channels. The paper considers α values of 1, 0.75, 0.5 and 0.25. A value of 1 keeps the baseline width, while 0.5 approximately halves the channels throughout the network [4].')
p('The depthwise cost scales with α, while the pointwise cost scales with α². Since pointwise convolutions account for most of the computation, the overall reduction in computation and parameters is approximately quadratic. It is not an exact rule for every layer. For example, at 224 × 224 input resolution, reducing α from 1 to 0.5 lowers the reported cost from 569 million to 149 million multiply-adds and the parameter count from 4.2 million to 1.3 million. ImageNet accuracy falls from 70.6% to 63.7% [4, Table 6].')
h('Resolution multiplier')
p('The resolution multiplier, ρ, reduces the input image size and the spatial size of the internal feature maps. The original paper tests input resolutions of 224, 192, 160 and 128. Computation scales approximately with ρ², while the number of learned parameters stays the same for a fixed width [4].')
p('For example, reducing the input from 224 × 224 to 128 × 128 at full width lowers computation from 569 million to 186 million multiply-adds. The parameter count remains 4.2 million, and ImageNet accuracy decreases from 70.6% to 64.4% [4, Table 7]. Width and resolution can therefore be adjusted to meet the available resources, but the resulting loss of accuracy needs to be checked for the intended task.')

new(); h('Training and face verification')
p('In the original paper, MobileNet models were trained in TensorFlow using RMSprop with asynchronous gradient descent. The authors used less regularization and data augmentation than for larger models, including little or no weight decay on the depthwise filters [4]. These settings describe the original MobileNet experiments.')
p('The paper also studies face embeddings by training MobileNet to reproduce the outputs of a pretrained FaceNet model. This process is called knowledge distillation. The training objective minimizes the squared difference between the two models’ outputs [4, Section 4.7]. However, the original paper does not provide an LFW result for these models. Its face-embedding results should therefore be kept separate from the LFW comparison.')
p('For an LFW result, Chen et al. provide a MobileNetV1 baseline in their MobileFaceNets paper [5]. This is still a MobileNetV1 model; it is not the proposed MobileFaceNet architecture. Their face preprocessing uses detected facial landmarks to align images to 112 × 112 pixels, with RGB values normalized by subtracting 127.5 and dividing by 128. For the baseline mobile networks, the first convolution uses stride 1 instead of stride 2 to preserve spatial detail [5].')
p('The baseline networks were trained from scratch on CASIA-WebFace using ArcFace loss. Training used stochastic gradient descent with momentum 0.9 and a batch size of 512. The learning rate started at 0.1 and was divided by 10 at 36,000, 52,000 and 58,000 iterations. Training ended at 60,000 iterations [5, Section 4.1]. This setup differs from both the original ImageNet training and the FaceNet distillation experiment.')
h('Reported performance')
p('Under this face-verification setup, MobileNetV1 achieves 98.63% accuracy on LFW and 88.95% on AgeDB-30, with 3.2 million parameters. Its reported inference time is 60 ms on a Qualcomm Snapdragon 820 CPU using four threads and the NCNN framework [5, Table 2]. This timing belongs to that device and implementation; it does not establish the speed of our complete webcam system.')
p('The 3.2 million parameters describe the face-verification baseline, while the 4.2 million figure from the original paper describes the ImageNet classifier. These should be recorded as different configurations. The supplied references do not report a YTF result for this MobileNetV1 baseline, and Table 2 does not specify its embedding dimension. The 128-dimensional embedding described for MobileFaceNet should not be assigned to MobileNetV1.')
h('Strengths and limitations')
p('MobileNet’s main strength is its low computational cost and the ability to adjust its width and input resolution. This gives us practical options when memory or processing power is limited. However, reducing either setting can lower accuracy, and a general-purpose mobile architecture is not automatically the best choice for faces. In the same CASIA-WebFace experiment, the face-specific MobileFaceNet reaches 99.28% on LFW with 0.99 million parameters and a 24 ms inference time [5, Table 2]. This shows why MobileNetV1 should be evaluated as a candidate rather than selected only because it is designed for mobile devices.')

new(); h('Comparison with the other architectures')
p('For our comparison, we use the LFW result, parameter count and inference time from the same MobileNetV1 face-verification experiment [5]. The ImageNet results explain the efficiency of the original architecture, but they describe a different configuration and task. Differences in training data, face alignment and evaluation settings also need to be considered when comparing MobileNetV1 with the other architectures. The reported inference time must be read together with the hardware and software used for testing.')
table=d.add_table(rows=1, cols=2)
table.autofit=False; table.columns[0].width=Inches(2.05); table.columns[1].width=Inches(4.75)
table.rows[0].cells[0].text='Comparison field'; table.rows[0].cells[1].text='MobileNetV1 face-verification baseline [5]'
rows=[('LFW accuracy','98.63%'),('YTF accuracy','Not reported'),('Parameters','3.2 million'),('Embedding dimension','Not explicitly reported for this baseline'),('Inference time','60 ms'),('Testing hardware and software','Snapdragon 820 CPU, four threads, NCNN'),('Training data and loss','CASIA-WebFace, ArcFace loss'),('Additional evaluation','88.95% on AgeDB-30')]
for a,b in rows:
    cells=table.add_row().cells; cells[0].text=a; cells[1].text=b
for i,row in enumerate(table.rows):
    trpr=row._tr.get_or_add_trPr(); trpr.append(OxmlElement('w:cantSplit'))
    if i==0: trpr.append(OxmlElement('w:tblHeader'))
    for cell in row.cells:
        tcpr=cell._tc.get_or_add_tcPr()
        borders=OxmlElement('w:tcBorders')
        for edge in ['top','left','bottom','right']:
            el=OxmlElement('w:'+edge); el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'4'); el.set(qn('w:color'),'D9D9D9'); borders.append(el)
        tcpr.append(borders)
        margins=OxmlElement('w:tcMar')
        for edge in ['top','left','bottom','right']:
            el=OxmlElement('w:'+edge); el.set(qn('w:w'),'90'); el.set(qn('w:type'),'dxa'); margins.append(el)
        tcpr.append(margins)
        if i==0:
            fill=OxmlElement('w:shd'); fill.set(qn('w:fill'),'E7E7E7'); tcpr.append(fill)
        for para in cell.paragraphs:
            para.paragraph_format.space_after=Pt(0)
            for run in para.runs: run.font.size=Pt(10.5)
h('Conclusion for MobileNetV1')
p('MobileNetV1 reduces computation by separating spatial filtering from channel combination. Its width and resolution multipliers allow the model to be adjusted to the resources available. The later face-verification baseline shows that it can achieve 98.63% accuracy on LFW with 3.2 million parameters [5]. This makes it a useful candidate for our system, although the reported result does not establish its performance on our own identities or hardware. We will need to measure recognition accuracy and the complete processing time before deciding whether it provides the balance required for real-time use.')
h('Conclusion for the main paper')
p('This study compares LightCNN-9, FaceNet, VGGFace and MobileNetV1 for a real-time facial recognition system. The reported results show different approaches to balancing recognition performance and computational cost. LightCNN-9 learns a compact facial representation, FaceNet learns distances between face embeddings, VGGFace uses a deep convolutional network, and MobileNetV1 reduces computation through depthwise separable convolutions. However, differences in training data, evaluation settings and hardware prevent a direct overall ranking. Our final choice will depend on testing the models on our own dataset and hardware, considering recognition accuracy, model size and the processing time of the complete system.')
h('References')
p('[4] A. G. Howard et al., “MobileNets: Efficient Convolutional Neural Networks for Mobile Vision Applications,” arXiv:1704.04861v1, 2017.')
p('[5] S. Chen, Y. Liu, X. Gao, and Z. Han, “MobileFaceNets: Efficient CNNs for Accurate Real-Time Face Verification on Mobile Devices,” arXiv:1804.07573v4, 2018.')
d.core_properties.title='Architecture 4 MobileNetV1'
d.core_properties.subject='MobileNetV1 architecture analysis and paper conclusion'
d.save(OUT/'Architecture_4_MobileNetV1.docx')
print(OUT/'Architecture_4_MobileNetV1.docx')
