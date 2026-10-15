from pathlib import Path
from docx import Document
from docx.oxml.ns import qn
from zipfile import ZipFile
from PIL import Image
import shutil

path = Path(r"C:\Users\bash\Desktop\vggface_gauri.docx")
doc = Document(path)

print(f"PARAGRAPHS={len(doc.paragraphs)} TABLES={len(doc.tables)} INLINE_SHAPES={len(doc.inline_shapes)} SECTIONS={len(doc.sections)}")
for i, sec in enumerate(doc.sections):
    cols = sec._sectPr.xpath('./w:cols')
    col_count = cols[0].get(qn('w:num')) if cols else None
    print(f"SECTION {i}: page={sec.page_width}x{sec.page_height} margins={sec.top_margin},{sec.right_margin},{sec.bottom_margin},{sec.left_margin} cols={col_count}")

for i, p in enumerate(doc.paragraphs):
    jc = p._p.xpath('./w:pPr/w:jc')
    align = jc[0].get(qn('w:val')) if jc else None
    fmt = p.paragraph_format
    drawings = len(p._p.xpath('.//w:drawing'))
    equations = len(p._p.xpath('.//m:oMath | .//m:oMathPara'))
    print(f"P{i:03d}|style={p.style.name!r}|align={align}|before={fmt.space_before}|after={fmt.space_after}|line={fmt.line_spacing}|drawings={drawings}|equations={equations}|text={p.text!r}")
    for j, r in enumerate(p.runs):
        if r.text or r._r.xpath('.//w:drawing'):
            color = r.font.color.rgb if r.font.color and r.font.color.type else None
            print(f"  R{j:02d}|text={r.text!r}|bold={r.bold}|italic={r.italic}|underline={r.underline}|size={r.font.size}|font={r.font.name}|color={color}|draw={len(r._r.xpath('.//w:drawing'))}")

for ti, table in enumerate(doc.tables):
    print(f"TABLE {ti}: rows={len(table.rows)} cols={len(table.columns)} style={table.style.name if table.style else None}")
    for ri, row in enumerate(table.rows):
        print(f"  ROW {ri}: " + " || ".join(cell.text.replace('\n', ' <NL> ') for cell in row.cells))

media_dir = Path(r"C:\Users\bash\Documents\GitHub\Dl_project_report_grp_zero_day\tmp\vggface_media")
media_dir.mkdir(parents=True, exist_ok=True)
with ZipFile(path) as zf:
    for name in zf.namelist():
        if name.startswith("word/media/") and not name.endswith("/"):
            out = media_dir / Path(name).name
            with zf.open(name) as src, out.open("wb") as dst:
                shutil.copyfileobj(src, dst)
            try:
                with Image.open(out) as im:
                    print(f"MEDIA {out.name}: {im.size[0]}x{im.size[1]} {im.format}")
            except Exception as exc:
                print(f"MEDIA {out.name}: unreadable {exc}")

for i, p in enumerate(doc.paragraphs):
    for blip in p._p.xpath('.//a:blip'):
        rid = blip.get(qn('r:embed'))
        rel = doc.part.rels[rid]
        extents = p._p.xpath('.//wp:extent')
        extent = (extents[0].get('cx'), extents[0].get('cy')) if extents else None
        print(f"DRAWING P{i:03d}: rid={rid} target={rel.target_ref} extent={extent}")
